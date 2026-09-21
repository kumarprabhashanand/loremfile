"""M3.5 office: generators, validators, negative tests, determinism tests.

The interesting one here is openpyxl. docs/06 §4 used to claim that these three
libraries were safe because their internal *part names* are deterministic. The spike
that preceded this milestone showed the claim was aimed at the wrong thing: openpyxl
spools each worksheet to a temporary file and adds it with ``ZipFile.write``, so that
one entry is stamped from the filesystem and no clock patch can reach it. Raw output
changed on six of eight consecutive runs. ``util.zipnorm`` is what actually makes xlsx
reproducible, and :func:`test_zipnorm_is_what_makes_xlsx_reproducible` is the negative
control that would notice if someone decided normalisation was cosmetic.
"""

from __future__ import annotations

import io
import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from loremfile import build as build_module
from loremfile.build import load_generators
from loremfile.catalog import Catalog, Fixture
from loremfile.generators import office
from loremfile.generators.base import REGISTRY, GeneratorContext
from loremfile.util import zipnorm
from loremfile.util.determinism import deterministic
from loremfile.validators import _REGISTRY, ValidationError, load, validate

WORKDIR = Path(tempfile.gettempdir())
CATALOG = Catalog.load()
load_generators()
load()

OOXML = ("docx/1page.docx", "xlsx/1sheet-10rows.xlsx", "pptx/1slide.pptx")
EVERY_OFFICE_FIXTURE = [
    path
    for path, entry in CATALOG.by_path.items()
    if entry.format in {"docx", "xlsx", "pptx", "rtf", "epub"}
]


def fixture(path: str) -> Fixture:
    return CATALOG.by_path[path]


def resolve_from_build(path: str) -> bytes:
    target = build_module.fixtures_dir() / path
    if not target.is_file():
        pytest.skip(f"run `loremfile build --only {path}` first")
    return target.read_bytes()


def build(path: str) -> bytes:
    entry = fixture(path)
    context = GeneratorContext(path=path, workdir=WORKDIR)
    context._dependency = resolve_from_build
    with deterministic(context.seed):
        out = REGISTRY.get(entry.generator)(context, **entry.params)
    assert isinstance(out, bytes)
    return out


def check(path: str) -> dict:
    entry = fixture(path)
    report = validate(build(path), entry, CATALOG.mime_for(entry))
    assert report.ok, report.failures
    return report.props


def measure(path: str, payload: bytes) -> dict:
    entry = fixture(path)
    return _REGISTRY[entry.format](payload, entry, CATALOG.mime_for(entry))


# --- every fixture validates ------------------------------------------------


@pytest.mark.parametrize("path", EVERY_OFFICE_FIXTURE)
def test_every_office_fixture_matches_its_catalog_entry(path: str) -> None:
    if fixture(path).size_class == "approx":
        pytest.skip("sized fixtures are covered by the build; generating one per test is slow")
    check(path)


# --- the determinism findings, each with a control --------------------------


def test_zipnorm_is_what_makes_xlsx_reproducible() -> None:
    """openpyxl's raw output carries a filesystem timestamp; the normalised one does not.

    Rather than sleep to let the clock move — which would make this test slow and still
    flaky — the raw archive is inspected directly: if any entry is stamped at something
    other than zipnorm's fixed epoch, raw bytes cannot be stable.
    """
    data = build("xlsx/1sheet-10rows.xlsx")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        stamps = {info.filename: info.date_time for info in archive.infolist()}
    assert set(stamps.values()) == {zipnorm.ZIP_EPOCH}, (
        "a normalised archive must carry exactly one timestamp; "
        f"found {sorted(set(stamps.values()))}"
    )

    # The control: the same workbook without normalisation is stamped from the clock,
    # which is what makes the normalisation load-bearing rather than tidy.
    workbook = office._new_workbook(GeneratorContext(path="xlsx/x.xlsx", workdir=WORKDIR))
    buffer = io.BytesIO()
    workbook.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        raw = {info.date_time for info in archive.infolist()}
    assert raw != {zipnorm.ZIP_EPOCH}, (
        "openpyxl now writes zipnorm's epoch by itself. If that is a real change in the "
        "pinned openpyxl, this test and docs/06 §4's row should both be revisited."
    )


@pytest.mark.parametrize("path", OOXML)
def test_core_properties_are_pinned_not_taken_from_the_clock(path: str) -> None:
    """The published bytes must not record when they happened to be built."""
    with zipfile.ZipFile(io.BytesIO(build(path))) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert "2020-01-01" in core, core
    assert office.AUTHOR in core


def test_epoch_is_built_at_call_time_so_the_guard_can_swap_the_class() -> None:
    """Inside the guard, `datetime.datetime` is a stand-in and libraries type-check it.

    A module-level constant would be an instance of the real class and python-docx would
    reject it with a message naming the same type twice.
    """
    context = GeneratorContext(path="docx/x.docx", workdir=WORKDIR)
    outside = office._epoch(context)
    with deterministic(context.seed):
        inside = office._epoch(context)
        assert type(inside) is not type(outside), (
            "the guard is no longer swapping datetime.datetime; _epoch's call-time "
            "construction can be simplified if that is now permanent"
        )
        assert isinstance(inside, __import__("datetime").datetime)
    assert inside == outside


def test_the_libraries_are_imported_at_module_scope() -> None:
    """Importing them inside the guard poisons them for the rest of the process."""
    source = Path(office.__file__).read_text(encoding="utf-8")
    body = source[source.index("def _epoch") :]
    for library in ("import docx", "import openpyxl", "from openpyxl", "from pptx"):
        assert library not in body, (
            f"{library!r} appears below the module header. Imported inside the "
            "determinism guard it binds the guard's datetime stand-in permanently."
        )


# --- structure, not parseability -------------------------------------------


def test_formula_cells_carry_cached_values() -> None:
    """A formula with no cached value reads as None in pandas and openpyxl alike."""
    data = build("xlsx/with-formulas.xlsx")
    values = load_workbook(io.BytesIO(data), data_only=True)["Orders"]
    formulas = load_workbook(io.BytesIO(data))["Orders"]
    assert formulas["G1"].value == "=SUM(B2:B11)"
    # quantities cycle 1..7, so ten rows are 1,2,3,4,5,6,7,1,2,3.
    assert values["G1"].value == 34
    assert values["G3"].value == "yes"
    assert values["D2"].value == round(1 * 2.5, 2)


def test_an_uncached_formula_is_rejected() -> None:
    """The negative control for the check above: openpyxl's own output must fail."""
    context = GeneratorContext(path="xlsx/with-formulas.xlsx", workdir=WORKDIR)
    with deterministic(context.seed):
        workbook = office._new_workbook(context)
        sheet = workbook.active
        sheet.title = "Orders"
        sheet.append(["item", "quantity"])
        sheet["C1"] = "=SUM(B2:B3)"
        uncached = office._save_workbook(workbook)
    with pytest.raises(ValidationError, match="no cached value"):
        measure("xlsx/with-formulas.xlsx", uncached)


def test_epub_mimetype_entry_is_first_and_stored() -> None:
    data = build("epub/epub3-3chapters.epub")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
    assert infos[0].filename == "mimetype"
    assert infos[0].compress_type == zipfile.ZIP_STORED


def test_epub_with_a_compressed_mimetype_is_rejected() -> None:
    """Readers reject this, and a plain zip check would not notice."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", b"application/epub+zip")
        archive.writestr("META-INF/container.xml", b"<container/>")
    with pytest.raises(ValidationError):
        measure("epub/epub3-3chapters.epub", out.getvalue())


def test_rtf_paragraph_count_does_not_confuse_par_with_pard() -> None:
    r"""`\pard` starts with `\par`; counting substrings doubled every paragraph."""
    props = check("rtf/simple.rtf")
    assert props["paragraphs"] == 3


def test_unbalanced_rtf_is_rejected() -> None:
    with pytest.raises(ValidationError, match="unclosed group"):
        measure("rtf/simple.rtf", b"{\\rtf1\\ansi{\\fonttbl}\nhello\\par\n")


def test_a_truncated_package_is_rejected() -> None:
    """Half a docx still starts with the zip magic bytes."""
    data = build("docx/1page.docx")
    with pytest.raises(ValidationError):
        measure("docx/1page.docx", data[: len(data) // 2])


def test_a_docx_without_the_document_part_is_rejected() -> None:
    """A zip full of the right-looking names is not a Word document."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("_rels/.rels", b"<Relationships/>")
    with pytest.raises(ValidationError, match="missing required parts"):
        measure("docx/1page.docx", out.getvalue())


# --- content ---------------------------------------------------------------


def test_the_table_holds_the_shared_dataset_with_a_header() -> None:
    props = check("docx/with-table.docx")
    assert props["table_rows"] == 11, "ten records plus a header row"


def test_images_come_from_the_published_fixtures() -> None:
    """The embedded PNG must be the published bytes, not a regenerated lookalike."""
    published = resolve_from_build("png/100x100.png")
    data = build("docx/with-images.docx")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        media = [archive.read(n) for n in archive.namelist() if n.startswith("word/media/")]
    assert published in media


def test_every_cell_type_survives_the_round_trip() -> None:
    data = build("xlsx/with-types-and-formats.xlsx")
    sheet = load_workbook(io.BytesIO(data))["Types"]
    labels = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert labels["boolean true"] is True
    assert labels["integer"] == 42
    assert labels["text that looks numeric"] == "007", "must stay text, not become 7"
    assert labels["empty"] is None
    assert str(labels["date"]).startswith("2020-01-01")


def test_the_people_sheet_is_a_prefix_of_the_bigger_one() -> None:
    """docs/05 §2's prefix property, at the spreadsheet level."""
    small = load_workbook(io.BytesIO(build("xlsx/1sheet-10rows.xlsx")))["People"]
    large = load_workbook(io.BytesIO(build("xlsx/1sheet-1000rows.xlsx")))["People"]
    for row in range(1, 12):
        assert [c.value for c in small[row]] == [c.value for c in large[row]]


def test_slides_are_widescreen() -> None:
    props = check("pptx/10slides.pptx")
    assert props["slides"] == 10
    assert props["slide_width_emu"] > props["slide_height_emu"]
