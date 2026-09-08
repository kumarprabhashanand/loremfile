"""Office and e-book fixtures: docx, xlsx, pptx, rtf, epub (docs/05 §3.5).

Four things had to be settled by spike before any of this could be catalogued, because
a nondeterministic fixture that reaches the manifest is frozen there permanently.

1. **openpyxl is not reproducible on its own.** It spools each worksheet to a temporary
   file and adds it with ``ZipFile.write`` (``openpyxl/writer/excel.py``), so that entry
   takes its timestamp from the filesystem — which no clock patch can reach. Raw output
   changed on six of eight consecutive runs; every entry *except* the worksheet sat at
   the patched epoch. ``zipnorm.normalize`` restamps everything and is what actually
   makes xlsx stable, so it is not optional here. python-docx and python-pptx build
   their archives in memory and were byte-identical without it; they are normalised too,
   for entry order and permissions.

2. **These libraries must be imported at module scope, never inside the guard.** They do
   ``from datetime import datetime`` at import time. Imported inside ``deterministic()``
   they bind the guard's stand-in class permanently, and then reject real datetimes once
   the guard exits — openpyxl raised ``TypeError: expected _FixedDatetime`` in the
   validator, in the same process, long after the generator had finished.

3. **Core properties are set explicitly**, as ``pdf.py`` does for fpdf2: relying on the
   patched clock is not enough when a library formats or type-checks the value itself.
   The epoch is built at call time by :func:`_epoch`, never stored as a constant — see
   that function for why.

4. **Cached formula values** are injected into the sheet XML afterwards. openpyxl has no
   API for them and already writes an empty ``<v/>`` placeholder, so the value has to
   replace that element; appending a second ``<v>`` leaves readers taking the empty one.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from typing import Any

import docx
from docx.shared import Inches
from lxml import etree
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image
from pptx import Presentation
from pptx.presentation import Presentation as PresentationDocument
from pptx.util import Inches as PptxInches
from pptx.util import Pt

from loremfile.config import APPROX_TOLERANCE
from loremfile.generators.base import GeneratorContext, generator
from loremfile.util import lorem, zipnorm
from loremfile.util.sizing import fit

#: Everything these documents claim about their own authorship.
AUTHOR = "loremfile.dev"
TITLE = "loremfile.dev sample document"

#: The spreadsheet namespace, for the cached-value injection.
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

#: Side of the seeded-noise PNG plates the sized fixtures embed (docs/05 §6).
PLATE_SIDE = 256

#: Measured size of one PLATE_SIDE plate plus its paragraph, used to pick the first `n`.
APPROX_PLATE_BYTES = 196_000

#: The threshold the IF formula tests, published in the fixture as a literal.
LARGE_ORDER = 10

#: Columns narrow enough to read in a portrait page.
TABLE_COLUMNS = ("id", "first_name", "last_name", "city", "country")


def _epoch(ctx: GeneratorContext) -> dt.datetime:
    """The fixed epoch, as an instance of whatever ``datetime.datetime`` now is.

    Built on every call rather than stored as a module constant. Inside the determinism
    guard ``datetime.datetime`` is a stand-in subclass, and python-docx checks its
    argument with ``isinstance(value, datetime.datetime)`` resolved through the module —
    so a real datetime built at import time is rejected with the genuinely confusing
    ``property requires <type 'datetime.datetime'> object, got <class 'datetime.datetime'>``.
    Building it here means the value is always the class currently installed.
    """
    return dt.datetime.fromtimestamp(ctx.epoch, tz=dt.UTC).replace(tzinfo=None)


def _plate(ctx: GeneratorContext, index: int, side: int = PLATE_SIDE) -> bytes:
    """One incompressible seeded-noise PNG, so a sized document is honestly sized."""
    needed = side * side * 3
    noise = ctx.stream(needed + index + 1)[index : index + needed]
    image = Image.frombytes("RGB", (side, side), noise)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=False, compress_level=6)
    return buffer.getvalue()


# --- docx ------------------------------------------------------------------


def _new_document(ctx: GeneratorContext) -> docx.document.Document:
    document = docx.Document()
    core = document.core_properties
    core.author = core.last_modified_by = AUTHOR
    core.title = TITLE
    core.created = core.modified = _epoch(ctx)
    core.revision = 1
    core.language = "en-GB"
    return document


def _save_docx(document: docx.document.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return zipnorm.normalize(buffer.getvalue())


@generator()
def docx_basic(ctx: GeneratorContext, *, paragraphs: int = 6, heading: bool = True) -> bytes:
    """Lorem Ipsum in the default Word styles, no embedded media."""
    document = _new_document(ctx)
    if heading:
        document.add_heading("Loremfile sample document", level=1)
    rng = ctx.rng
    for _ in range(paragraphs):
        document.add_paragraph(lorem.paragraph(rng))
    return _save_docx(document)


@generator()
def docx_with_images(ctx: GeneratorContext, *, images: list[str]) -> bytes:
    """Embeds published image fixtures, read through ``ctx.dependency``."""
    document = _new_document(ctx)
    document.add_heading("Images", level=1)
    rng = ctx.rng
    for path in images:
        document.add_paragraph(lorem.paragraph(rng))
        document.add_picture(io.BytesIO(ctx.dependency(path)), width=Inches(3.0))
    return _save_docx(document)


@generator()
def docx_with_table(ctx: GeneratorContext, *, rows: int = 10) -> bytes:
    """One table of the shared people dataset, with a header row."""
    document = _new_document(ctx)
    document.add_heading("People", level=1)
    people = ctx.dataset("people", rows)
    table = document.add_table(rows=1, cols=len(TABLE_COLUMNS))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, TABLE_COLUMNS, strict=True):
        cell.text = column
        cell.paragraphs[0].runs[0].font.bold = True
    for person in people:
        cells = table.add_row().cells
        for cell, column in zip(cells, TABLE_COLUMNS, strict=True):
            cell.text = str(person[column])
    return _save_docx(document)


@generator(parallel_safe=False)
def docx_sized(ctx: GeneratorContext, *, size: int) -> bytes:
    """One lorem paragraph per embedded noise plate, fitted to ``size`` (docs/05 §6)."""

    def make(plates: int) -> bytes:
        document = _new_document(ctx)
        document.add_heading("Loremfile sized document", level=1)
        rng = ctx.rng
        for index in range(plates):
            document.add_paragraph(lorem.paragraph(rng))
            document.add_picture(io.BytesIO(_plate(ctx, index)), width=Inches(3.0))
        return _save_docx(document)

    start = max(1, round(size / APPROX_PLATE_BYTES))
    _, payload = fit(size, APPROX_TOLERANCE, make, n0=start, n_min=1, n_max=1000)
    return payload


# --- xlsx ------------------------------------------------------------------


def _new_workbook(ctx: GeneratorContext) -> Workbook:
    workbook = Workbook()
    properties = workbook.properties
    properties.creator = properties.lastModifiedBy = AUTHOR
    properties.title = TITLE
    properties.created = properties.modified = _epoch(ctx)
    return workbook


def _save_workbook(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    # Not cosmetic: openpyxl stamps worksheet entries from a temp file's mtime.
    return zipnorm.normalize(buffer.getvalue())


def _cell_value(value: object) -> object:
    """Coerce a dataset value into something Excel can hold, losing as little as possible.

    The shared people dataset carries three things a worksheet cell cannot take: a list
    (`tags`), a timezone-aware datetime (`signup_at` — Excel has no concept of one) and
    a Decimal. The list is joined, the timezone dropped after conversion to UTC, and the
    Decimal left alone, which openpyxl writes as a number. Anything else is passed
    through, so a new dataset column fails loudly here rather than being silently
    stringified.
    """
    if isinstance(value, list):
        return ",".join(str(item) for item in value)
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        return value.astimezone(dt.UTC).replace(tzinfo=None)
    return value


def _write_people(sheet: Worksheet, people: list[dict[str, Any]]) -> None:
    columns = list(people[0])
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for person in people:
        sheet.append([_cell_value(person[column]) for column in columns])
    for column, width in (("D", 30), ("F", 24), ("J", 20)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"


@generator(parallel_safe=False)
def xlsx_people(ctx: GeneratorContext, *, rows: int, sheets: int = 1) -> bytes:
    """The shared people dataset, one row per record."""
    workbook = _new_workbook(ctx)
    first = workbook.active
    if first is None:  # pragma: no cover - openpyxl always creates one
        raise RuntimeError("the new workbook has no active sheet")
    first.title = "People" if sheets == 1 else "People 1"
    _write_people(first, ctx.dataset("people", rows))
    for index in range(2, sheets + 1):
        extra = workbook.create_sheet(f"People {index}")
        _write_people(extra, ctx.dataset("people", rows))
    return _save_workbook(workbook)


def _inject_cached_values(data: bytes, sheet: str, cached: dict[str, object]) -> bytes:
    """Give every formula cell the value a spreadsheet application would cache.

    Without this, `load_workbook(data_only=True)` and pandas both read `None`: openpyxl
    writes the formula and an empty `<v/>`, and never computes anything. The value has to
    replace that placeholder — appending a second `<v>` leaves readers taking the first.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts[sheet])
    written: set[str] = set()
    for cell in root.iter(f"{{{SHEET_NS}}}c"):
        reference = cell.get("r")
        if reference not in cached:
            continue
        value = cached[reference]
        if isinstance(value, str):
            cell.set("t", "str")
        element = cell.find(f"{{{SHEET_NS}}}v")
        if element is None:
            element = etree.SubElement(cell, f"{{{SHEET_NS}}}v")
        element.text = str(value)
        written.add(reference)
    missing = sorted(set(cached) - written)
    if missing:
        raise ValueError(f"no formula cells at {missing} to cache values into")
    parts[sheet] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for name, payload in parts.items():
            target.writestr(name, payload)
    return zipnorm.normalize(out.getvalue())


@generator()
def xlsx_with_formulas(ctx: GeneratorContext, *, rows: int = 10) -> bytes:
    """SUM, AVERAGE, IF and VLOOKUP over an order table, with cached results."""
    workbook = _new_workbook(ctx)
    sheet = workbook.active
    if sheet is None:  # pragma: no cover
        raise RuntimeError("the new workbook has no active sheet")
    sheet.title = "Orders"
    sheet.append(["item", "quantity", "unit_price", "total"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    quantities = [(index % 7) + 1 for index in range(rows)]
    prices = [round(2.5 + index * 1.25, 2) for index in range(rows)]
    for index in range(rows):
        row = index + 2
        sheet.append([f"item-{index + 1:03d}", quantities[index], prices[index]])
        sheet.cell(row=row, column=4).value = f"=B{row}*C{row}"

    last = rows + 1
    totals = [round(q * p, 2) for q, p in zip(quantities, prices, strict=True)]
    grand = round(sum(totals), 2)
    mean = round(grand / rows, 10)
    lookup_row = min(3, rows)

    sheet["F1"] = "quantity total"
    sheet["F2"] = "mean line total"
    sheet["F3"] = "is a large order"
    sheet["F4"] = f"unit price of item-{lookup_row:03d}"
    sheet["G1"] = f"=SUM(B2:B{last})"
    sheet["G2"] = f"=AVERAGE(D2:D{last})"
    sheet["G3"] = f'=IF(G1>{LARGE_ORDER},"yes","no")'
    sheet["G4"] = f'=VLOOKUP("item-{lookup_row:03d}",A2:C{last},3,FALSE)'
    sheet.column_dimensions["F"].width = 26

    cached: dict[str, object] = {
        "G1": sum(quantities),
        "G2": mean,
        "G3": "yes" if sum(quantities) > LARGE_ORDER else "no",
        "G4": prices[lookup_row - 1],
    }
    for index in range(rows):
        cached[f"D{index + 2}"] = totals[index]
    return _inject_cached_values(_save_workbook(workbook), "xl/worksheets/sheet1.xml", cached)


@generator()
def xlsx_with_types(ctx: GeneratorContext) -> bytes:
    """One column per Excel type and number format, including the awkward ones."""
    workbook = _new_workbook(ctx)
    sheet = workbook.active
    if sheet is None:  # pragma: no cover
        raise RuntimeError("the new workbook has no active sheet")
    sheet.title = "Types"
    moment = _epoch(ctx)

    rows: list[tuple[str, Any, str | None]] = [
        ("text", "Lorem ipsum", None),
        ("integer", 42, "0"),
        ("float", 3.14159, "0.00000"),
        ("negative", -1234.5, "#,##0.00"),
        ("boolean true", True, None),
        ("boolean false", False, None),
        ("date", moment.date(), "yyyy-mm-dd"),
        ("time", dt.time(13, 45, 30), "hh:mm:ss"),
        ("datetime", moment, "yyyy-mm-dd hh:mm:ss"),
        ("percentage", 0.4275, "0.00%"),
        ("currency", 1999.99, '"£"#,##0.00'),
        ("scientific", 0.000000123, "0.00E+00"),
        ("text that looks numeric", "007", "@"),
        ("text that looks like a date", "2020-01-01", "@"),
        ("leading apostrophe kept", "+44 20 7946 0000", "@"),
        ("empty", None, None),
    ]
    sheet.append(["label", "value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for label, value, number_format in rows:
        sheet.append([label, value])
        cell = sheet.cell(row=sheet.max_row, column=2)
        if number_format:
            cell.number_format = number_format
        cell.alignment = Alignment(horizontal="left")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 24
    return _save_workbook(workbook)


@generator(parallel_safe=False)
def xlsx_sized(ctx: GeneratorContext, *, size: int) -> bytes:
    """One people row per embedded noise plate, fitted to ``size`` (docs/05 §6)."""

    def make(plates: int) -> bytes:
        workbook = _new_workbook(ctx)
        sheet = workbook.active
        if sheet is None:  # pragma: no cover
            raise RuntimeError("the new workbook has no active sheet")
        sheet.title = "People"
        _write_people(sheet, ctx.dataset("people", max(plates, 1)))
        images = workbook.create_sheet("Plates")
        for index in range(plates):
            picture = SheetImage(io.BytesIO(_plate(ctx, index)))
            images.add_image(picture, f"A{index * 14 + 1}")
        return _save_workbook(workbook)

    start = max(1, round(size / APPROX_PLATE_BYTES))
    _, payload = fit(size, APPROX_TOLERANCE, make, n0=start, n_min=1, n_max=1000)
    return payload


# --- pptx ------------------------------------------------------------------


#: python-pptx's built-in template is 4:3 (10 x 7.5 inches). Every fixture here is
#: widescreen, which is what current PowerPoint and Keynote produce; the 4:3 shape gets
#: its own named fixture (`4x3-5slides.pptx`, P1b) rather than being the silent default.
SLIDE_16_9 = (PptxInches(13.333), PptxInches(7.5))


def _new_presentation(ctx: GeneratorContext) -> PresentationDocument:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = SLIDE_16_9
    core = presentation.core_properties
    core.author = core.last_modified_by = AUTHOR
    core.title = TITLE
    core.created = core.modified = _epoch(ctx)
    core.revision = 1
    return presentation


def _save_presentation(presentation: PresentationDocument) -> bytes:
    buffer = io.BytesIO()
    presentation.save(buffer)
    return zipnorm.normalize(buffer.getvalue())


def _title_and_body(presentation: PresentationDocument, title: str, body: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    frame = slide.placeholders[1].text_frame
    frame.text = body
    frame.paragraphs[0].font.size = Pt(18)


@generator()
def pptx_basic(ctx: GeneratorContext, *, slides: int = 1) -> bytes:
    """Title-and-content slides of Lorem Ipsum, 16:9."""
    presentation = _new_presentation(ctx)
    rng = ctx.rng
    for index in range(slides):
        _title_and_body(presentation, f"Slide {index + 1}", lorem.paragraph(rng))
    return _save_presentation(presentation)


@generator(parallel_safe=False)
def pptx_sized(ctx: GeneratorContext, *, size: int) -> bytes:
    """One slide per embedded noise plate, fitted to ``size`` (docs/05 §6)."""

    def make(plates: int) -> bytes:
        presentation = _new_presentation(ctx)
        rng = ctx.rng
        for index in range(plates):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = f"Plate {index + 1}"
            slide.shapes.add_picture(
                io.BytesIO(_plate(ctx, index)),
                PptxInches(1),
                PptxInches(1.8),
                PptxInches(3),
                PptxInches(3),
            )
            notes = slide.notes_slide.notes_text_frame
            notes.text = lorem.sentence(rng)
        return _save_presentation(presentation)

    start = max(1, round(size / APPROX_PLATE_BYTES))
    _, payload = fit(size, APPROX_TOLERANCE, make, n0=start, n_min=1, n_max=1000)
    return payload


# --- rtf -------------------------------------------------------------------

#: RTF is a plain-text format, so these fixtures are written by hand rather than by a
#: library. Every byte is ASCII: non-ASCII characters would need \u escapes, and these
#: two fixtures deliberately stay in the range every reader agrees on.
RTF_HEADER = r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\froman Times New Roman;}}"


def _rtf_escape(text: str) -> str:
    """Escape RTF's three special characters, and refuse anything non-ASCII."""
    if not text.isascii():
        raise ValueError(f"RTF fixtures stay in ASCII; got {text!r}")
    return text.replace("\\", r"\\").replace("{", r"\{").replace("}", r"\}")


@generator()
def rtf_simple(ctx: GeneratorContext, *, paragraphs: int = 3) -> bytes:
    """Plain paragraphs of Lorem Ipsum, hand-written RTF."""
    rng = ctx.rng
    body = [r"\pard\sa200\sl276\slmult1\f0\fs24 " + _rtf_escape(lorem.paragraph(rng)) + r"\par"]
    for _ in range(paragraphs - 1):
        body.append(r"\pard\sa200 " + _rtf_escape(lorem.paragraph(rng)) + r"\par")
    return (RTF_HEADER + "\n" + "\n".join(body) + "\n}").encode("ascii")


# --- epub ------------------------------------------------------------------

#: EPUB 3 needs these three files plus the content documents (OCF 3.0, EPUB 3.3).
EPUB_CONTAINER = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="EPUB/package.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""

EPUB_CSS = """body { font-family: serif; line-height: 1.5; margin: 1em; }
h1 { font-size: 1.4em; }
"""


def _epub_chapter(number: int, title: str, paragraphs: list[str]) -> str:
    body = "\n".join(f"    <p>{paragraph}</p>" for paragraph in paragraphs)
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"
      xml:lang="en" lang="en">
  <head>
    <title>{title}</title>
    <meta charset="utf-8"/>
    <link rel="stylesheet" type="text/css" href="style.css"/>
  </head>
  <body>
    <section epub:type="chapter" id="chapter-{number}">
    <h1>{title}</h1>
{body}
    </section>
  </body>
</html>
"""


@generator()
def epub3(ctx: GeneratorContext, *, chapters: int = 3, paragraphs: int = 4) -> bytes:
    """A valid EPUB 3 with a package document, a nav document and N chapters.

    The identifier and the modification date are fixed rather than generated: an EPUB
    carries a `dcterms:modified` that would otherwise be the wall clock, and a UUID that
    would otherwise be random. Both are part of the published bytes forever.
    """
    rng = ctx.rng
    moment = _epoch(ctx).strftime("%Y-%m-%dT%H:%M:%SZ")
    identifier = "urn:uuid:" + ctx.stream(16).hex()

    titles = [f"Chapter {number}" for number in range(1, chapters + 1)]
    files: dict[str, bytes] = {
        "mimetype": b"application/epub+zip",
        "META-INF/container.xml": EPUB_CONTAINER.encode("utf-8"),
        "EPUB/style.css": EPUB_CSS.encode("utf-8"),
    }
    for number, title in enumerate(titles, start=1):
        text = [lorem.paragraph(rng) for _ in range(paragraphs)]
        files[f"EPUB/chapter-{number}.xhtml"] = _epub_chapter(number, title, text).encode("utf-8")

    items = "\n".join(
        f'    <item id="chapter-{number}" href="chapter-{number}.xhtml" '
        f'media-type="application/xhtml+xml"/>'
        for number in range(1, chapters + 1)
    )
    spine = "\n".join(
        f'    <itemref idref="chapter-{number}"/>' for number in range(1, chapters + 1)
    )
    files["EPUB/package.opf"] = f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="pub-id"
         xml:lang="en">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="pub-id">{identifier}</dc:identifier>
    <dc:title>Loremfile Sample Book</dc:title>
    <dc:language>en</dc:language>
    <dc:creator>{AUTHOR}</dc:creator>
    <dc:rights>CC0 1.0 Universal</dc:rights>
    <meta property="dcterms:modified">{moment}</meta>
  </metadata>
  <manifest>
    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
    <item id="css" href="style.css" media-type="text/css"/>
{items}
  </manifest>
  <spine>
{spine}
  </spine>
</package>
""".encode()

    links = "\n".join(
        f'        <li><a href="chapter-{number}.xhtml">{title}</a></li>'
        for number, title in enumerate(titles, start=1)
    )
    files["EPUB/nav.xhtml"] = f"""<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"
      xml:lang="en" lang="en">
  <head><title>Contents</title><meta charset="utf-8"/></head>
  <body>
    <nav epub:type="toc" id="toc">
      <h1>Contents</h1>
      <ol>
{links}
      </ol>
    </nav>
  </body>
</html>
""".encode()

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)
    # epub=True puts `mimetype` first and stores it uncompressed, which OCF requires.
    return zipnorm.normalize(out.getvalue(), epub=True)
