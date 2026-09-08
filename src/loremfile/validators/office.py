"""Validators for docx, xlsx, pptx, rtf and epub (docs/06 §6 and §13).

Two independent readers, as docs/06 §6 requires. The writing library is never the only
opinion: every OOXML fixture is also read as a **zip archive** and checked against what
the format actually mandates — `[Content_Types].xml` present, the package relationships
present, the media parts really there. A file python-docx happily reopens can still be a
package Word refuses, because python-docx reads back its own conventions.

These validators assert **structure, not parseability** (docs/06 §13): that a table has
the rows it claims, that a formula cell carries a cached value a non-calculating reader
can use, that an EPUB's `mimetype` entry is first and stored uncompressed. Any of those
can be wrong in a file that opens without complaint.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

import docx
from openpyxl import load_workbook
from pptx import Presentation

from loremfile.catalog import Fixture
from loremfile.validators import ValidationError, register
from loremfile.validators.text import text_props

#: Parts every OOXML package must carry, whatever the flavour.
OOXML_REQUIRED = ("[Content_Types].xml", "_rels/.rels")

#: A table that is only a header row, or only data, is a mistake in the generator.
MIN_TABLE_ROWS = 2

#: EPUB's Open Container Format requires exactly this, and readers rely on it.
EPUB_MIMETYPE = b"application/epub+zip"


def _archive(data: bytes) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValidationError(f"not a readable zip container: {exc}") from exc


def _check_package(data: bytes, required: tuple[str, ...]) -> list[str]:
    """The zip-level reading: the parts an OOXML consumer looks for before parsing."""
    with _archive(data) as archive:
        names = archive.namelist()
        broken = archive.testzip()
        if broken is not None:
            raise ValidationError(f"corrupt entry in the package: {broken}")
    missing = [name for name in (*OOXML_REQUIRED, *required) if name not in names]
    if missing:
        raise ValidationError(f"package is missing required parts: {missing}")
    return names


@register("docx")
def validate_docx(data: bytes, fixture: Fixture, _mime: str) -> dict[str, Any]:
    names = _check_package(data, ("word/document.xml",))
    try:
        document = docx.Document(io.BytesIO(data))
        paragraphs = [p for p in document.paragraphs if p.text.strip()]
        tables = document.tables
        table_rows = [len(table.rows) for table in tables]
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"python-docx could not read it: {exc}") from exc

    images = [name for name in names if name.startswith("word/media/")]
    if fixture.params.get("images") and not images:
        raise ValidationError("declares images but the package has no word/media parts")
    for index, rows in enumerate(table_rows):
        if rows < MIN_TABLE_ROWS:
            raise ValidationError(f"table {index} has {rows} row(s): no header plus data")
    return {
        "paragraphs": len(paragraphs),
        "tables": len(tables),
        "table_rows": table_rows[0] if table_rows else 0,
        "images": len(images),
        "words": sum(len(p.text.split()) for p in paragraphs),
    }


@register("xlsx")
def validate_xlsx(data: bytes, _fixture: Fixture, _mime: str) -> dict[str, Any]:
    _check_package(data, ("xl/workbook.xml",))
    try:
        workbook = load_workbook(io.BytesIO(data))
        cached = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValidationError(f"openpyxl could not read it: {exc}") from exc

    sheet = workbook.worksheets[0]
    formulas = [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if formulas:
        # The whole point of caching values: a reader that does not calculate must still
        # see numbers. openpyxl writes an empty <v/> placeholder, which reads back as
        # None — so an uncached formula fixture is useless without ever failing to parse.
        values = cached[sheet.title]
        empty = [cell.coordinate for cell in formulas if values[cell.coordinate].value is None]
        if empty:
            raise ValidationError(f"formula cells carry no cached value: {sorted(empty)[:5]}")

    rows = sheet.max_row
    if rows < 1:
        raise ValidationError("the first worksheet has no rows")
    return {
        "sheets": len(workbook.worksheets),
        "sheet_names": ",".join(workbook.sheetnames),
        "rows": rows,
        "columns": sheet.max_column,
        "formulas": len(formulas),
        "images": sum(len(worksheet._images) for worksheet in workbook.worksheets),
    }


@register("pptx")
def validate_pptx(data: bytes, _fixture: Fixture, _mime: str) -> dict[str, Any]:
    names = _check_package(data, ("ppt/presentation.xml",))
    try:
        presentation = Presentation(io.BytesIO(data))
        slides = list(presentation.slides)
    except Exception as exc:
        raise ValidationError(f"python-pptx could not read it: {exc}") from exc

    if not slides:
        raise ValidationError("has no slides")
    layouts = {slide.slide_layout.name for slide in slides}
    return {
        "slides": len(slides),
        "layouts": len(layouts),
        "images": len([name for name in names if name.startswith("ppt/media/")]),
        "slide_width_emu": presentation.slide_width,
        "slide_height_emu": presentation.slide_height,
        "notes": sum(1 for slide in slides if slide.has_notes_slide),
    }


@register("rtf")
def validate_rtf(data: bytes, _fixture: Fixture, mime: str) -> dict[str, Any]:
    """RTF has no library reader in the toolchain, so the structure is checked directly.

    Brace balance is the check that matters: an unbalanced RTF opens as an empty document
    in some readers and as garbage in others, and neither reports an error.
    """
    if not data.startswith(b"{\\rtf1"):
        raise ValidationError("does not start with the RTF signature {\\rtf1")
    try:
        text = data.decode("ascii")
    except UnicodeDecodeError as exc:
        raise ValidationError(f"contains non-ASCII bytes at {exc.start}") from exc

    depth = 0
    for index, character in enumerate(text):
        if index and text[index - 1] == "\\":
            continue
        if character == "{":
            depth += 1
        elif character == "}":
            depth -= 1
            if depth < 0:
                raise ValidationError(f"closing brace with no opening one at byte {index}")
    if depth != 0:
        raise ValidationError(f"{depth} unclosed group(s)")
    return {
        # `\par` is a prefix of `\pard`, so a plain count double-counts every paragraph
        # that opens with a paragraph-default reset. RTF control words end at the first
        # non-letter, which is what the lookahead encodes.
        **text_props(data, mime),
        "paragraphs": len(re.findall(r"\\par(?![a-zA-Z])", text)),
        "has_font_table": r"\fonttbl" in text,
    }


@register("epub")
def validate_epub(data: bytes, _fixture: Fixture, _mime: str) -> dict[str, Any]:
    """EPUB is a zip with rules the zip format itself will not enforce."""
    with _archive(data) as archive:
        names = archive.namelist()
        infos = archive.infolist()
        if names[0] != "mimetype":
            raise ValidationError(f"first entry is '{names[0]}', not 'mimetype' (OCF 3.0 §4.1)")
        if infos[0].compress_type != zipfile.ZIP_STORED:
            raise ValidationError("the mimetype entry is compressed; OCF requires it stored")
        if archive.read("mimetype") != EPUB_MIMETYPE:
            raise ValidationError("the mimetype entry does not hold application/epub+zip")
        if "META-INF/container.xml" not in names:
            raise ValidationError("no META-INF/container.xml")
        container = archive.read("META-INF/container.xml").decode("utf-8")
        opf = [name for name in names if name.endswith(".opf")]
        if not opf:
            raise ValidationError("no package document (.opf) in the container")
        package = archive.read(opf[0]).decode("utf-8")
        chapters = [name for name in names if name.endswith(".xhtml")]

    if opf[0] not in container:
        raise ValidationError(f"container.xml does not point at the package document {opf[0]}")
    for element in ("<dc:identifier", "<dc:title", "<dc:language", "dcterms:modified"):
        if element not in package:
            raise ValidationError(f"the package document has no {element}")
    if 'properties="nav"' not in package:
        raise ValidationError("the package document declares no nav document (EPUB 3 §5.2)")
    return {
        "chapters": len([name for name in chapters if "chapter" in name]),
        "documents": len(chapters),
        "epub_version": "3.0" if 'version="3.0"' in package else "2.0",
        "has_nav": True,
    }
