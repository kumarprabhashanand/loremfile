"""Plain text from a fixture, through a real reader for each format (docs/04 §1.3.2).

Used to check `contains_text`, never to publish: readers differ in whitespace, line breaks
and the order of blocks, so what one of them extracts is right for that reader only. Each
function here is a common library's ordinary text path, not a parser written for us.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable

import docx
import html5lib
from markdown_it import MarkdownIt
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

#: Elements whose text a reader does not show.
HIDDEN = frozenset({"script", "style", "template", "head"})


def normalize(text: str) -> str:
    """Every run of whitespace as one space: the form `contains_text` is compared in."""
    return " ".join(text.split())


def _pdf(data: bytes, password: str | None) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt(password or "")
    return "\n".join(page.extract_text() for page in reader.pages)


def _docx(data: bytes, _password: str | None) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def _pptx(data: bytes, _password: str | None) -> str:
    presentation = Presentation(io.BytesIO(data))
    return "\n".join(
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def _xlsx(data: bytes, _password: str | None) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return "\n".join(
            str(value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        )
    finally:
        workbook.close()


def _html_text(markup: bytes | str) -> str:
    root = html5lib.parse(markup, namespaceHTMLElements=False)
    parts: list[str] = []

    def walk(element: object) -> None:
        tag = getattr(element, "tag", None)
        if not isinstance(tag, str) or tag in HIDDEN:  # a comment's tag is a function
            return
        if getattr(element, "text", None):
            parts.append(element.text)  # type: ignore[attr-defined]
        for child in element:  # type: ignore[attr-defined]
            walk(child)
            if child.tail:
                parts.append(child.tail)

    walk(root)
    return "".join(parts)


def _html(data: bytes, _password: str | None) -> str:
    return _html_text(data)


def _markdown(data: bytes, _password: str | None) -> str:
    return _html_text(MarkdownIt("commonmark").render(data.decode("utf-8")))


def _epub(data: bytes, _password: str | None) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        documents = sorted(name for name in archive.namelist() if name.endswith(".xhtml"))
        return "\n".join(_html_text(archive.read(name)) for name in documents)


#: Format → reader. A format absent here cannot carry `contains_text`.
READERS: dict[str, Callable[[bytes, str | None], str]] = {
    "pdf": _pdf,
    "docx": _docx,
    "pptx": _pptx,
    "xlsx": _xlsx,
    "html": _html,
    "md": _markdown,
    "epub": _epub,
}


def text(fmt: str, data: bytes, *, password: str | None = None) -> str:
    """The text a reader of ``fmt`` extracts, whitespace normalized."""
    return normalize(READERS[fmt](data, password))
